"""Build an editable, content-preserving cheatsheet from reviewed JSON."""
from pathlib import Path
from copy import deepcopy
import argparse
import hashlib
import json
import re
from lxml import etree
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
ROOT = Path(__file__).resolve().parents[1]
CJK = re.compile(r'([\u2e80-\ua4cf\uac00-\ud7af\uf900-\ufaff\ufe30-\ufe4f\uff01-\uff60\u3000-\u303f]+)')
FORMULA = re.compile(r'\[\[F:([\w.-]+)\]\]')
COLORS = {'core': 'FFF0A8', 'normal': 'DAECFF', 'case': 'E0F2DD'}

def read_data(path):
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    assert isinstance(data.get('title'), str) and data['title'].strip(), 'Missing title'
    assert isinstance(data.get('sections'), dict) and data['sections'], 'Missing sections'
    assert isinstance(data.get('items'), list) and data['items'], 'Missing items'
    ids = []
    for item in data['items']:
        for key in ('id', 'section', 'question', 'answer'):
            assert isinstance(item.get(key), str) and item[key].strip(), f'Missing {key}'
        assert not re.search(r'\s', item['id']), 'IDs must not contain whitespace'
        assert item['section'] in data['sections'], f'Undefined section: {item["section"]}'
        assert item.get('priority', 'normal') in COLORS, 'Unknown priority'
        assert '[[MATH' not in item['answer'], 'Convert extraction placeholders to reviewed formulas'
        for name in FORMULA.findall(item['answer']):
            assert name in data.get('formulas', {}), f'Undefined formula: {name}'
            formula_text(data['formulas'][name])
        ids.append(item['id'])
    assert len(ids) == len(set(ids)), 'Duplicate question IDs'
    return data

def formula_text(formula):
    if 'runs' in formula:
        result = ''.join(x['text'] for x in formula['runs'])
        assert result == formula.get('plain_text', result), 'Formula text does not match its runs'
        assert all(not (x.get('superscript') and x.get('subscript')) for x in formula['runs'])
        return result
    assert isinstance(formula.get('omml'), str), 'Formula needs runs or OMML'
    assert isinstance(formula.get('plain_text'), str), 'OMML needs a reviewed plain_text equivalent'
    return formula['plain_text']

def answer_text(item, data):
    return FORMULA.sub(lambda m: formula_text(data['formulas'][m.group(1)]), item['answer'])

def run_text(p, text, size, *, color='000000', fill=None, bold=False, superscript=False, subscript=False):
    for chunk in CJK.split(text):
        if not chunk:
            continue
        run = p.add_run(chunk)
        family = 'Arial Unicode MS' if CJK.fullmatch(chunk) else 'Arial Narrow'
        run.font.name = family
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.cs_bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if superscript:
            run.font.superscript = True
        if subscript:
            run.font.subscript = True
        rp = run._element.get_or_add_rPr()
        for key in ('ascii', 'hAnsi', 'cs'):
            rp.rFonts.set(qn('w:' + key), family)
        rp.rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'en-US')
        lang.set(qn('w:eastAsia'), 'zh-CN')
        rp.append(lang)
        if fill:
            el = OxmlElement('w:shd')
            el.set(qn('w:val'), 'clear')
            el.set(qn('w:fill'), fill)
            rp.append(el)

def paragraph(p, font, line, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(line)
    pf.keep_with_next = keep_next
    pf.keep_together = True
    pf.widow_control = False
    prop = p._p.get_or_add_pPr()
    mark = OxmlElement('w:rPr')
    for key in ('sz', 'szCs'):
        e = OxmlElement('w:' + key)
        e.set(qn('w:val'), str(round(font * 2)))
        mark.append(e)
    prop.append(mark)
    for key, value in [('snapToGrid', '0'), ('autoSpaceDE', '0'), ('autoSpaceDN', '0'), ('suppressAutoHyphens', '1')]:
        e = OxmlElement('w:' + key)
        e.set(qn('w:val'), value)
        prop.append(e)

def write_answer(p, item, data, font):
    term_colors = {}
    for key, color in [('core_terms', 'FFF0A8'), ('case_terms', 'E0F2DD'), ('logic_terms', 'FFDCDD')]:
        term_colors.update({term: color for term in data.get(key, []) if term})
    literal = '|'.join(re.escape(x) for x in sorted(term_colors, key=len, reverse=True))
    numeric = r'\d+(?:[,.]\d+)*(?:%|\s*(?:months?|years?|million|billion|HKD|HK\$|USD))?'
    pattern = re.compile('(?:' + literal + ')|' + numeric if literal else numeric)
    def plain(text):
        offset = 0
        for m in pattern.finditer(text):
            run_text(p, text[offset:m.start()], font)
            run_text(p, m.group(), font, fill=term_colors.get(m.group(), 'DAECFF'))
            offset = m.end()
        run_text(p, text[offset:], font)
    offset = 0
    for m in FORMULA.finditer(item['answer']):
        plain(item['answer'][offset:m.start()])
        f = data['formulas'][m.group(1)]
        if 'runs' in f:
            for r in f['runs']:
                run_text(p, r['text'], font, fill='DAECFF', superscript=r.get('superscript', False), subscript=r.get('subscript', False))
        else:
            el = etree.fromstring(f['omml'].encode(), parser=etree.XMLParser(resolve_entities=False, no_network=True))
            assert el.tag in (qn('m:oMath'), qn('m:oMathPara')), 'Expected native OMML root'
            p._p.append(deepcopy(el))
        offset = m.end()
    plain(item['answer'][offset:])

def build(data, output, font=4, line=None, columns=8, margin=4, overwrite=False):
    output = Path(output)
    assert overwrite or not output.exists(), 'Existing output: use --overwrite only for an intended replacement'
    assert 1 <= columns <= 12 and font > 0 and abs(font * 2 - round(font * 2)) < 1e-8
    line = font + 0.3 if line is None else line
    assert line >= font and margin >= 0
    doc = Document(ROOT / 'assets/reference.docx')
    for child in list(doc._element.body):
        if child.tag != qn('w:sectPr'):
            doc._element.body.remove(child)
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Mm(297), Mm(210)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(margin)
    s.header_distance = s.footer_distance = Mm(0)
    for ref in s._sectPr.xpath('./w:headerReference|./w:footerReference'):
        s._sectPr.remove(ref)
    cols = s._sectPr.find(qn('w:cols'))
    cols.set(qn('w:num'), str(columns))
    cols.set(qn('w:space'), '57')
    cols.set(qn('w:sep'), '1')
    for child in list(cols):
        cols.remove(child)
    grid = s._sectPr.find(qn('w:docGrid'))
    if grid is not None:
        s._sectPr.remove(grid)
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Arial Narrow', Pt(font)
    style.paragraph_format.line_spacing = Pt(line)
    defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    for key in ('sz', 'szCs'):
        e = defaults.find(qn('w:' + key))
        if e is None:
            e = OxmlElement('w:' + key)
            defaults.append(e)
        e.set(qn('w:val'), str(round(font * 2)))
    p = doc.add_paragraph(style='Title')
    paragraph(p, font, line + 0.8, True)
    run_text(p, data['title'], font + 0.5, color='1F4E79')
    previous = None
    for item in data['items']:
        if item['section'] != previous:
            p = doc.add_paragraph()
            paragraph(p, font, line + 0.8, True)
            p.paragraph_format.space_before = Pt(1.3)
            run_text(p, data['sections'][item['section']], font + 0.5, color='1F4E79', fill='DAECFF')
            previous = item['section']
        p = doc.add_paragraph()
        paragraph(p, font, line)
        run_text(p, item['id'] + ' ', font, color='B00020', fill=COLORS[item.get('priority', 'normal')])
        run_text(p, item['question'], font, color='B00020')
        run_text(p, ' A: ', font, bold=True)
        write_answer(p, item, data, font)
    doc.core_properties.title = data['title']
    doc.core_properties.subject = 'Two-sided A4 question-and-answer cheatsheet'
    doc.core_properties.author = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {'items': len(data['items']), 'font_pt': font, 'line_pt': line, 'columns': columns,
            'margin_mm': margin, 'output': str(output.resolve()), 'sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
            'render_status': 'not yet rendered; page count must be verified'}

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input')
    parser.add_argument('--output', required=True)
    parser.add_argument('--font', type=float, default=4)
    parser.add_argument('--line', type=float)
    parser.add_argument('--columns', type=int, default=8)
    parser.add_argument('--margin', type=float, default=4)
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    print(json.dumps(build(read_data(args.input), args.output, args.font, args.line, args.columns, args.margin, args.overwrite), ensure_ascii=False))
